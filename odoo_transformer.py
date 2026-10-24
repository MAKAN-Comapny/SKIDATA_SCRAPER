"""
Copy of SKIDATA-Odoo/Transformer.py so the scraper can build the Odoo file itself.
Keep the two in step if the Odoo column mapping changes.

Transform Parking Transactions export into Odoo-ready invoice lines.

Steps:
  1. Load Parking-Transactions .xlsx from exports/
  2. Keep rows where Payment is not 0 / blank
  3. Map/rename columns to the Odoo structure
  4. Write result to Cleaned-Exports/ and clear exports/
"""
import os
import glob
import math
import re
from decimal import Decimal, InvalidOperation
from typing import Optional

import pandas as pd
from openpyxl import Workbook

EXPORTS_DIR = "exports"
CLEANED_EXPORTS_DIR = "Cleaned-Exports"
PARKING_FILE_MARKER = "Parking-Transactions"

# Exact output column names (Odoo import structure)
OUT_PARTNER = "Partner"
OUT_INVOICE_BILL_DATE = "Invoice/Bill Date"
OUT_INVOICE_LABEL = "Invoice Lines/Label"
OUT_QUANTITY = "Invoice Lines/Quantity"
OUT_UNIT_PRICE = "Invoice Lines/Unit Price"
OUT_TAXES = "invoice Lines/Taxes"
OUT_IN_TIME = "In Time"
OUT_OUT_TIME = "Out Time"
OUT_TICKET_SEQUENCE = "Ticket Sequence"
OUT_PAYMENT_REFERENCE = "Payment Reference"
OUT_ANALYTICS = "Invoice Lines/Analytics Distribution"
OUT_PLATE_NUMBER = "Plate Number"
OUT_LOCATION = "Location"
OUT_REGION = "Region"
OUT_SERVICE = "Service"
OUT_SUB_SERVICE = "Sub Service"

FIXED_PARTNER = "SKIDATA"
FIXED_QUANTITY = 1
FIXED_TAXES = "15%"
ANALYTICS_DISTRIBUTION_PERCENT = 100

VIEW_MALL_CAR_PARK = "View Mall"
VIEW_MALL_PREMIUM_RATE = 15  # SAR per hour
VIEW_MALL_PREMIUM_MIN_DURATION_MINUTES = 15

# SKIDATA Car Park name -> Odoo Location / Region / Service / Sub Service
CAR_PARK_MAPPING = {
    "Westfield DAMMAM": {
        "location": "02 - Nakheel Mall Dammam",
        "region": "Dammam Region",
        "service": "Paid Parking",
        "sub_service": "Standard",
    },
    "View Mall": {
        "location": "01 - The View Mall Riyadh",
        "region": "Riyadh Region",
        "service": "Paid Parking",
        "sub_service": "Standard",
    },
    "Alnoor Mall": {
        "location": "06 - Al Noor Mall Madinah",
        "region": "Madinah Region",
        "service": "Paid Parking",
        "sub_service": "Standard",
    },
    "Mall of Arabia": {
        "location": "03 - Arab Mall Jeddah",
        "region": "Jeddah Region",
        "service": "Paid Parking",
        "sub_service": "Standard",
    },
    "Premium Carpark": {
        "location": "06 - Al Noor Mall Madinah",
        "region": "Madinah Region",
        "service": "Paid Parking",
        "sub_service": "Premium",
    },
}

LOCATION_ANALYTIC_IDS = {
    "06 - Al Noor Mall Madinah": 54,
    "03 - Arab Mall Jeddah": 46,
    "02 - Nakheel Mall Dammam": 43,
    "01 - The View Mall Riyadh": 42,
}

REGION_ANALYTIC_IDS = {
    "Dammam Region": 111,
    "Jeddah Region": 112,
    "Madinah Region": 115,
    "Riyadh Region": 110,
}

SERVICE_ANALYTIC_IDS = {
    "Paid Parking": 135,
    "Valet Parking": 121,
}

SUB_SERVICE_ANALYTIC_IDS = {
    "Premium": 131,
    "Standard": 130,
}

OUTPUT_COLUMNS = [
    OUT_PARTNER,
    OUT_INVOICE_BILL_DATE,
    OUT_INVOICE_LABEL,
    OUT_QUANTITY,
    OUT_UNIT_PRICE,
    OUT_TAXES,
    OUT_IN_TIME,
    OUT_OUT_TIME,
    OUT_TICKET_SEQUENCE,
    OUT_PAYMENT_REFERENCE,
    OUT_ANALYTICS,
    OUT_PLATE_NUMBER,
    OUT_LOCATION,
    OUT_REGION,
    OUT_SERVICE,
    OUT_SUB_SERVICE,
]


def _normalize_header(name: str) -> str:
    """Lowercase and collapse separators for fuzzy column matching."""
    return re.sub(r"[\s/_-]+", "", str(name).strip().lower())


def _find_column(columns, *candidates: str) -> str:
    """Return the actual column name matching any candidate."""
    normalized = {_normalize_header(c): c for c in columns}
    for candidate in candidates:
        key = _normalize_header(candidate)
        if key in normalized:
            return normalized[key]
    raise KeyError(
        f"Could not find column matching any of {candidates}. "
        f"Available columns: {list(columns)}"
    )


def _is_blank(value) -> bool:
    """Treat empty / NaN / SKIDATA plate placeholders (-, #) as missing."""
    if value is None:
        return True
    if isinstance(value, float) and pd.isna(value):
        return True
    text = str(value).strip()
    return text == "" or text.lower() == "nan" or text in ("-", "#")


def _pick_plate(entry_plate, exit_plate) -> str:
    """Use whichever plate exists; if both exist take either; else blank."""
    entry_blank = _is_blank(entry_plate)
    exit_blank = _is_blank(exit_plate)
    if not entry_blank:
        return str(entry_plate).strip()
    if not exit_blank:
        return str(exit_plate).strip()
    return ""


def _serial_as_text(value) -> str:
    """Keep full serial/ticket sequence as text (never as Excel number)."""
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        as_int = int(round(value))
        if abs(value - as_int) < 1e-9:
            return str(as_int)
        return format(value, "f").rstrip("0").rstrip(".")
    text = str(value).strip()
    if text == "" or text.lower() == "nan":
        return ""
    if re.fullmatch(r"[+-]?\d+(\.\d+)?[eE][+-]?\d+", text):
        try:
            d = Decimal(text)
            if d == d.to_integral_value():
                return format(d.to_integral_value(), "f")
            return format(d, "f")
        except InvalidOperation:
            return text
    if re.fullmatch(r"[+-]?\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def _payment_is_nonzero(series: pd.Series) -> pd.Series:
    """True where Payment has a numeric value other than 0."""
    numeric = pd.to_numeric(series, errors="coerce")
    return numeric.notna() & (numeric != 0)


def _map_car_park(car_park_value) -> dict:
    """Map SKIDATA Car Park to Odoo Location / Region / Service / Sub Service."""
    key = "" if car_park_value is None or (isinstance(car_park_value, float) and pd.isna(car_park_value)) else str(car_park_value).strip()
    mapped = CAR_PARK_MAPPING.get(key)
    if mapped is None:
        return {
            "location": "",
            "region": "",
            "service": "",
            "sub_service": "",
        }
    return mapped


def _actual_duration_minutes(entry_time, exit_time) -> Optional[float]:
    """Parking duration in minutes from Entry/Exit timestamps."""
    entry = pd.to_datetime(entry_time, errors="coerce")
    exit_ = pd.to_datetime(exit_time, errors="coerce")
    if pd.isna(entry) or pd.isna(exit_):
        return None
    return (exit_ - entry).total_seconds() / 60.0


def _paid_parking_hours(actual_duration_minutes: float) -> int:
    """Power BI: ROUNDUP(ROUNDUP(Actual_Duration, 0) / 60, 0)."""
    rounded_minutes = math.ceil(actual_duration_minutes)
    return math.ceil(rounded_minutes / 60.0)


def _is_view_mall_premium(car_park, entry_time, exit_time, payment) -> bool:
    """
    View Mall premium rule (Power BI):
      Paid Parking Hours = ROUNDUP(ROUNDUP(Actual_Duration, 0) / 60, 0)
      Premium = Actual_Duration > 15 AND Payment = Hours * 15 AND Payment <> 0
    """
    if str(car_park).strip() != VIEW_MALL_CAR_PARK:
        return False

    payment_num = pd.to_numeric(payment, errors="coerce")
    if pd.isna(payment_num) or payment_num == 0:
        return False

    duration = _actual_duration_minutes(entry_time, exit_time)
    if duration is None or duration <= VIEW_MALL_PREMIUM_MIN_DURATION_MINUTES:
        return False

    hours = _paid_parking_hours(duration)
    expected_payment = hours * VIEW_MALL_PREMIUM_RATE
    return float(payment_num) == float(expected_payment)


def _resolve_sub_service(car_park, entry_time, exit_time, payment, mapped_sub_service: str) -> str:
    """Override View Mall Sub Service to Premium when the rate rule matches."""
    if _is_view_mall_premium(car_park, entry_time, exit_time, payment):
        return "Premium"
    return mapped_sub_service


def _build_analytics_distribution(
    location: str, region: str, service: str, sub_service: str
) -> str:
    """
    Build Odoo analytics JSON: {"location_id":100, "region_id":100, ...}
    Example: {"54":100, "115":100, "135":100, "131":100}
    """
    ids = []
    for value, lookup in (
        (location, LOCATION_ANALYTIC_IDS),
        (region, REGION_ANALYTIC_IDS),
        (service, SERVICE_ANALYTIC_IDS),
        (sub_service, SUB_SERVICE_ANALYTIC_IDS),
    ):
        analytic_id = lookup.get(str(value).strip()) if value else None
        if analytic_id is not None:
            ids.append(analytic_id)

    parts = [f'"{analytic_id}":{ANALYTICS_DISTRIBUTION_PERCENT}' for analytic_id in ids]
    return "{" + ", ".join(parts) + "}"


def _as_date_only(value):
    """Convert timestamp to a date (no time component)."""
    ts = pd.to_datetime(value, errors="coerce")
    if pd.isna(ts):
        return None
    return ts.date()


def _write_odoo_excel(df: pd.DataFrame, dest_path: str) -> None:
    """Write workbook with Ticket Sequence forced as Excel text."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    headers = list(df.columns)
    ws.append(headers)
    ticket_col_idx = headers.index(OUT_TICKET_SEQUENCE) + 1  # 1-based
    payment_ref_col_idx = headers.index(OUT_PAYMENT_REFERENCE) + 1
    bill_date_col_idx = headers.index(OUT_INVOICE_BILL_DATE) + 1

    for row in df.itertuples(index=False, name=None):
        values = []
        for col_name, value in zip(headers, row):
            if col_name in (OUT_TICKET_SEQUENCE, OUT_PAYMENT_REFERENCE):
                values.append(_serial_as_text(value))
            elif pd.isna(value):
                values.append(None)
            else:
                values.append(value)
        ws.append(values)

    for row_idx in range(2, ws.max_row + 1):
        for col_idx in (ticket_col_idx, payment_ref_col_idx):
            text_cell = ws.cell(row=row_idx, column=col_idx)
            if text_cell.value is not None:
                text_cell.value = str(text_cell.value)
                text_cell.number_format = "@"

        date_cell = ws.cell(row=row_idx, column=bill_date_col_idx)
        if date_cell.value is not None:
            date_cell.number_format = "YYYY-MM-DD"

    os.makedirs(os.path.dirname(dest_path) or ".", exist_ok=True)
    wb.save(dest_path)


def transform_parking_file(src_path: str, dest_path: str) -> tuple[int, int]:
    """
    Filter + remap one Parking Transactions workbook.
    Returns (rows_before_filter, rows_after_filter).
    """
    # Resolve text columns, then re-read them as text to preserve digits.
    header_df = pd.read_excel(src_path, nrows=0)
    serial_col = _find_column(header_df.columns, "Serial Number", "SerialNumber")
    receipt_col = _find_column(header_df.columns, "Receipt Number", "ReceiptNumber")
    df = pd.read_excel(src_path, dtype={serial_col: str, receipt_col: str})
    rows_before = len(df)

    payment_col = _find_column(df.columns, "Payment")
    txn_date_col = _find_column(
        df.columns,
        "Transaction Date/Time",
        "Transaction Date / Time",
        "Transaction Date Time",
    )
    article_col = _find_column(df.columns, "Article")
    entry_date_col = _find_column(
        df.columns, "Entry Date / Time", "Entry Date/Time", "Entry Date Time"
    )
    exit_date_col = _find_column(
        df.columns, "Exit Date / Time", "Exit Date/Time", "Exit Date Time"
    )
    entry_plate_col = _find_column(
        df.columns, "Entry Plate Number", "Entry Plate", "EntryPlateNumber"
    )
    exit_plate_col = _find_column(
        df.columns, "Exit Plate Number", "Exit Plate", "ExitPlateNumber"
    )
    car_park_col = _find_column(df.columns, "Car Park", "CarPark")

    filtered = df[_payment_is_nonzero(df[payment_col])].copy()
    rows_after = len(filtered)

    location_maps = [_map_car_park(v) for v in filtered[car_park_col].values]
    car_parks = filtered[car_park_col].values
    payments = filtered[payment_col].values
    entry_times = filtered[entry_date_col].values
    exit_times = filtered[exit_date_col].values

    locations = [m["location"] for m in location_maps]
    regions = [m["region"] for m in location_maps]
    services = [m["service"] for m in location_maps]
    sub_services = [
        _resolve_sub_service(cp, ent, ext, pay, m["sub_service"])
        for cp, ent, ext, pay, m in zip(
            car_parks, entry_times, exit_times, payments, location_maps
        )
    ]
    analytics = [
        _build_analytics_distribution(loc, reg, svc, sub)
        for loc, reg, svc, sub in zip(locations, regions, services, sub_services)
    ]

    out = pd.DataFrame(
        {
            OUT_PARTNER: FIXED_PARTNER,
            OUT_INVOICE_BILL_DATE: [
                _as_date_only(v) for v in filtered[txn_date_col].values
            ],
            OUT_INVOICE_LABEL: filtered[article_col].values,
            OUT_QUANTITY: FIXED_QUANTITY,
            OUT_UNIT_PRICE: pd.to_numeric(filtered[payment_col], errors="coerce").values,
            OUT_TAXES: FIXED_TAXES,
            OUT_IN_TIME: entry_times,
            OUT_OUT_TIME: exit_times,
            OUT_TICKET_SEQUENCE: [
                _serial_as_text(v) for v in filtered[serial_col].values
            ],
            OUT_PAYMENT_REFERENCE: [
                _serial_as_text(v) for v in filtered[receipt_col].values
            ],
            OUT_ANALYTICS: analytics,
            OUT_PLATE_NUMBER: [
                _pick_plate(e, x)
                for e, x in zip(
                    filtered[entry_plate_col].values,
                    filtered[exit_plate_col].values,
                )
            ],
            OUT_LOCATION: locations,
            OUT_REGION: regions,
            OUT_SERVICE: services,
            OUT_SUB_SERVICE: sub_services,
        }
    )
    out = out[OUTPUT_COLUMNS]
    _write_odoo_excel(out, dest_path)
    return rows_before, rows_after


def run():
    """Transform Parking Transactions files into Odoo format under Cleaned-Exports/."""
    all_files = glob.glob(os.path.join(EXPORTS_DIR, "*.xlsx"))
    parking_files = [f for f in all_files if PARKING_FILE_MARKER in os.path.basename(f)]

    if not parking_files:
        raise SystemExit(
            f"No Parking Transactions .xlsx files found in '{EXPORTS_DIR}'. "
            "Run scraper first."
        )

    os.makedirs(CLEANED_EXPORTS_DIR, exist_ok=True)

    for src in parking_files:
        base = os.path.basename(src)
        if base.lower().endswith(".xlsx"):
            dest_name = base[:-5] + "-Odoo.xlsx"
        else:
            dest_name = base + "-Odoo.xlsx"
        dest = os.path.join(CLEANED_EXPORTS_DIR, dest_name)

        before, after = transform_parking_file(src, dest)
        print(f"  {base}")
        print(f"    rows: {before} -> {after} (Payment != 0)")
        print(f"    wrote: {dest_name}")

    for file in all_files:
        os.remove(file)

    print(
        f"Transformed {len(parking_files)} Parking Transactions file(s) "
        f"to {CLEANED_EXPORTS_DIR}/"
    )


if __name__ == "__main__":
    run()
